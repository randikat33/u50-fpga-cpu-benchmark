#!/usr/bin/env python3
"""Build Supplementary File S1 (literature audit) as an Excel workbook."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
wb = Workbook()

readme = wb.active
readme.title = "README"
lines = [
    "Supplementary File S1 - Methodology audit of FPGA-vs-CPU comparisons (2018-2025)",
    "",
    "Inclusion: peer-reviewed paper (conference, workshop or journal), 2018-2025, full text openly available,",
    "reports at least one FPGA-vs-CPU speedup or energy ratio on a named CPU.",
    "Search: web search on FPGA + CPU + (speedup|energy|Alveo|HBM|OpenCL|HLS), then backward/forward snowballing",
    "from Table 1 of the paper. Open-access bias: papers without open full text were excluded.",
    "Coding: one coder; each code is backed by a verbatim quote (column 'evidence'). Codes are assigned from the quote,",
    "not from a summary. A second coder should re-check a random subset (planned before submission).",
    "",
    "Codebook (Y = stated/done, partial = partly, N = not done, NS = not stated):",
    "C1_mt   CPU baseline multithreaded (OpenMP/MPI/threads)",
    "C1_simd CPU baseline uses SIMD or a tuned library (AVX, NEON, MKL, AES-NI ...) and says so",
    "C2      number of CPU cores/threads/sockets stated",
    "C3      FPGA time includes / excludes / separately reports host-device transfers, or not stated (NS)",
    "C4      how power/energy was obtained (measured, estimated, TDP) or not reported",
    "C4b     measurement boundary stated (which components; idle handling)",
    "C5      CPU frequency, turbo or warm-up controlled or reported",
    "C6      run-to-run variability reported (std, CI, min/max) - averages alone count as 'partial'",
    "C7      equivalence of FPGA and CPU outputs checked",
    "C8      code or artefact available",
    "",
    "Sheets: coding (30 papers), excluded (screened out with reason), summary (percentages used in Section 2.5).",
]
for i, t in enumerate(lines, 1):
    readme.cell(row=i, column=1, value=t)
readme["A1"].font = Font(bold=True, size=12)
readme.column_dimensions["A"].width = 120


def sheet_from_csv(name, path, widths=None):
    ws = wb.create_sheet(name)
    rows = list(csv.reader(open(path)))
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "C2"
    for col in ws.columns:
        L = col[0].column_letter
        m = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[L].width = min(max(10, m + 2), (widths or {}).get(col[0].value, 45))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


sheet_from_csv("coding", os.path.join(HERE, "coding.csv"), {"evidence": 110, "source_url": 40})
sheet_from_csv("excluded", os.path.join(HERE, "excluded.csv"))
sheet_from_csv("summary", os.path.join(HERE, "audit_summary.csv"))
out = os.path.join(HERE, "S1_literature_audit.xlsx")
wb.save(out)
print("wrote", out)
