#!/usr/bin/env python3
"""
make_second_coder_sheet.py - blind coding workbook for the second coder of the methodology audit.

Writes S1_second_coder_blank.xlsx: the 30 included papers in a random order (seed 2026) with links, one
drop-down code column and one quote column per item, the codebook and instructions. The first coder's codes
and quotes are NOT in the workbook. Compare afterwards with audit_kappa.py.
"""
import csv, os, random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
rows = list(csv.DictReader(open(os.path.join(HERE, "coding.csv"))))

# item column -> (allowed codes, short definition shown in the header note)
ITEMS = [
    ("c1_mt", ["Y", "partial", "N", "NS"], "C1a CPU baseline multithreaded (OpenMP, MPI, threads, a threaded library)"),
    ("c1_simd", ["Y", "partial", "N", "NS"], "C1b CPU baseline uses SIMD or a tuned library (AVX, NEON, MKL, AES-NI, ...) AND the paper says so"),
    ("c2_cores", ["Y", "N"], "C2 number of CPU cores, threads or sockets used is stated"),
    ("c3_transfers", ["included", "excluded", "reported separately", "NS"], "C3 host-card transfers in the FPGA time"),
    ("c4_energy", ["measured both", "FPGA measured, CPU estimated/TDP", "estimated both", "not reported"], "C4 how power/energy was obtained"),
    ("c4b_boundary", ["Y", "partial", "N"], "C4b measurement boundary fully stated (which components are counted; idle handling)"),
    ("c5_freq_warmup", ["Y", "partial", "N"], "C5 CPU frequency, turbo or warm-up controlled or reported"),
    ("c6_variability", ["Y", "partial", "N"], "C6 run-to-run variability reported (std, CI, min-max; averages alone = partial)"),
    ("c7_equivalence", ["Y", "partial", "N"], "C7 equivalence of FPGA and CPU outputs checked"),
    ("c8_artifact", ["Y", "partial", "N"], "C8 code or artefact available"),
]

CODEBOOK = [
    ("General", "Code only what the paper (main text, appendix, or a supplement it links) states. Do not infer from "
                "other papers of the same group. Every code must be supported by a verbatim quote (copy-paste, with page "
                "or section if possible). If nothing is stated, the code is NS or N as listed, and the quote cell says 'none'."),
    ("Y / partial / N / NS", "Y = stated and done. partial = done for part of the results, or stated vaguely. "
                             "N = stated as not done, or clearly absent where it would have to be stated. NS = not stated."),
    ("C1a multithreaded", "Y if the CPU baseline used more than one thread/core (OpenMP, MPI, pthreads, a threaded library "
                          "such as MKL with threads). partial if only some CPU results are multithreaded. N if stated "
                          "single-threaded. NS if the paper does not say."),
    ("C1b SIMD / tuned library", "Y if the paper says the CPU code uses SIMD (AVX, SSE, NEON, auto-vectorisation reported) or a "
                                 "tuned library (MKL, OpenBLAS, FFTW, OpenSSL/AES-NI, oneDNN, a vendor-optimised DB engine ...). "
                                 "partial if only for part of the baseline. Compiler flags alone (-O3) are NS unless vectorisation is stated."),
    ("C2 cores stated", "Y if the number of CPU cores, threads or sockets used for the reported CPU numbers is given. "
                        "The CPU model alone is N."),
    ("C3 transfers", "included = FPGA time contains host-card data transfers; excluded = kernel time only; reported "
                     "separately = both are given; NS = cannot be told from the text (also for platforms where it is unclear)."),
    ("C4 energy", "measured both = power or energy measured on the FPGA and on the CPU (sensors, RAPL, meters, board APIs). "
                  "FPGA measured, CPU estimated/TDP = one side measured, the other from TDP or a model. estimated both = both "
                  "from tools/models (e.g. XPE, TDP). not reported = no power or energy numbers."),
    ("C4b boundary", "Only if C4 is not 'not reported'; otherwise N. Y if the paper says exactly which components are "
                     "counted on both sides (e.g. 'CPU package + DRAM via RAPL; card power via XRT, host excluded') and how "
                     "idle power is treated. partial if only some of this is stated."),
    ("C5 CPU clock", "Y if the CPU frequency governor, turbo state, fixed frequency, or warm-up runs are controlled or "
                     "reported. partial if only mentioned (e.g. nominal frequency) without control."),
    ("C6 variability", "Y if run-to-run variability is reported (standard deviation, confidence interval, min-max, error "
                       "bars with a stated meaning). partial if repetitions are averaged but no spread is given."),
    ("C7 equivalence", "Y if the paper states that FPGA and CPU outputs were compared and agree (bit-exact, within a "
                       "tolerance, same accuracy on a dataset). partial if only for some cases."),
    ("C8 artefact", "Y if code or an artefact is publicly available (a link that is given in the paper). partial if "
                    "only part (e.g. the FPGA kernel but not the CPU baseline) or 'available on request'."),
]


def main():
    order = [r for r in rows]
    random.Random(2026).shuffle(order)
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    text = [
        "Second-coder sheet for the methodology audit (JSA paper, Section 2.5; Supplementary File S1)",
        "",
        "Why: the audit was coded by one author. Reviewers asked for a second, independent coder and an agreement",
        "statistic (Cohen's kappa). Your codes are compared item by item with the first coder's (audit_kappa.py);",
        "disagreements are then discussed and resolved, and the paper reports kappa before resolution.",
        "",
        "Rules:",
        " 1. Work BLIND: do not open S1_literature_audit.xlsx or coding.csv until you have finished.",
        " 2. Code the papers in the order of the sheet 'coding' (random order). Open the link in column E.",
        " 3. For every item choose a code from the drop-down list and paste the supporting quote next to it.",
        " 4. Read the sheet 'codebook' first; it has the decision rules. When unsure, choose the more",
        "    conservative code and write a note in column 'notes'.",
        " 5. Record how long each paper took (minutes) in 'minutes'. Expect 20-30 minutes per paper.",
        " 6. If time is short: code at least the first 12 papers of the sheet (a random 40% subset).",
        " 7. Save as S1_second_coder_<your initials>.xlsx and send it back.",
        "",
        "Coder name: (fill in on the sheet 'coding', cell B1)",
    ]
    for t in text:
        ws.append([t])
    ws.column_dimensions["A"].width = 110

    cb = wb.create_sheet("codebook")
    cb.append(["Item", "Rule"])
    for a, b in CODEBOOK:
        cb.append([a, b])
    cb.column_dimensions["A"].width = 26
    cb.column_dimensions["B"].width = 120
    for row in cb.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(bold=True)

    c = wb.create_sheet("coding")
    c["A1"], c["B1"] = "Coder:", ""
    head = ["order", "id", "paper", "year", "link"]
    for k, _, _ in ITEMS:
        head += [k, k + "_quote"]
    head += ["minutes", "notes"]
    c.append([])
    c.append(head)
    hdr_row = 3
    for i, r in enumerate(order, 1):
        c.append([i, r["id"], f"{r['short']} ({r['venue']})", r["year"], r["source_url"]] + [""] * (2 * len(ITEMS) + 2))
        c.cell(row=hdr_row + i, column=5).hyperlink = r["source_url"]
    fill = PatternFill("solid", fgColor="DDEBF7")
    for j, h in enumerate(head, 1):
        cell = c.cell(row=hdr_row, column=j)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = hdr_row + len(order)
    for n, (k, codes, desc) in enumerate(ITEMS):
        col = 6 + 2 * n
        letter = c.cell(row=hdr_row, column=col).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(codes) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="code", error="choose a code from the list")
        dv.promptTitle = k
        dv.prompt = desc[:250]
        dv.showInputMessage = True
        c.add_data_validation(dv)
        dv.add(f"{letter}{hdr_row + 1}:{letter}{last}")
        c.column_dimensions[letter].width = 16
        c.column_dimensions[c.cell(row=hdr_row, column=col + 1).column_letter].width = 40
    for letter, w in (("A", 6), ("B", 7), ("C", 34), ("D", 6), ("E", 30)):
        c.column_dimensions[letter].width = w
    c.freeze_panes = "F4"
    out = os.path.join(HERE, "S1_second_coder_blank.xlsx")
    wb.save(out)
    print("wrote", out, "with", len(order), "papers; first 12 (random subset):", [r["id"] for r in order[:12]])


if __name__ == "__main__":
    main()
