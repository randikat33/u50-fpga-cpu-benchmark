#!/usr/bin/env python3
"""
audit_kappa.py - agreement between the two coders of the methodology audit.

Usage:  python3 audit_kappa.py S1_second_coder_<initials>.xlsx   (or a CSV with the same columns)

Compares the second coder's codes with coding.csv (first coder) for every paper the second coder finished.
For each item it reports, on the binary variable the paper counts (Table 3; e.g. 'Y' vs everything else):
  n, percent agreement, Cohen's kappa with a bootstrap 95% CI over papers (2,000 resamples, seed 1),
  PABAK (prevalence- and bias-adjusted kappa = 2 p_o - 1, useful when an item is rare, e.g. C5), and the
  count under each coder; plus Cohen's kappa on the full code set (Y/partial/N/NS; C3 and C4 categories).
Writes kappa_summary.csv, kappa_summary.json and disagreements.csv (both codes and both quotes side by side,
for the reconciliation meeting). Kappa is reported BEFORE reconciliation.
"""
import csv, json, os, sys
import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
ITEMS = ["c1_mt", "c1_simd", "c2_cores", "c3_transfers", "c4_energy", "c4b_boundary", "c5_freq_warmup",
         "c6_variability", "c7_equivalence", "c8_artifact"]


def c3_class(v):
    v = (v or "").strip().lower()
    for k in ("reported separately", "included", "excluded"):
        if v.startswith(k):
            return k
    return "NS" if v.startswith("ns") or not v else v


def c4_class(v):
    v = (v or "").strip()
    if not v or v.startswith("not"):
        return "none"
    if v == "measured both" or (v.startswith("measured") and "estimated" not in v and "TDP" not in v):
        return "measured"
    return "mixed"          # one side estimated/TDP, or both estimated


def full(item, v):
    if item == "c3_transfers":
        return c3_class(v)
    if item == "c4_energy":
        return c4_class(v)
    return (v or "").strip()


# binary variables as counted in the paper (Table 3)
BINARY = [
    ("C1a multithreaded", "c1_mt", lambda v: v == "Y", None),
    ("C1b SIMD/tuned library stated", "c1_simd", lambda v: v == "Y", None),
    ("C2 cores stated", "c2_cores", lambda v: v == "Y", None),
    ("C3 transfer treatment stated", "c3_transfers", lambda v: c3_class(v) != "NS", None),
    ("C3 transfers excluded", "c3_transfers", lambda v: c3_class(v) == "excluded", None),
    ("C4 power/energy reported", "c4_energy", lambda v: c4_class(v) != "none", None),
    ("C4 measured on both sides", "c4_energy", lambda v: c4_class(v) == "measured", None),
    ("C4b boundary fully stated (energy papers)", "c4b_boundary", lambda v: v == "Y", "energy"),
    ("C5 CPU clock/warm-up controlled", "c5_freq_warmup", lambda v: v == "Y", None),
    ("C6 variability reported", "c6_variability", lambda v: v == "Y", None),
    ("C7 output equivalence checked", "c7_equivalence", lambda v: v == "Y", None),
    ("C8 code/artefact available", "c8_artifact", lambda v: v == "Y", None),
]


def kappa(a, b):
    a, b = list(a), list(b)
    n = len(a)
    if n == 0:
        return np.nan, np.nan
    cats = sorted(set(a) | set(b))
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    k = (po - pe) / (1 - pe) if pe < 1 else np.nan      # no variation at all: kappa undefined
    return k, po


def boot_kappa(a, b, n=2000, seed=1):
    rng = np.random.default_rng(seed)
    a, b = np.asarray(a, dtype=object), np.asarray(b, dtype=object)
    ks = []
    for _ in range(n):
        i = rng.integers(0, len(a), len(a))
        k, _ = kappa(a[i], b[i])
        if np.isfinite(k):
            ks.append(k)
    return (float(np.percentile(ks, 2.5)), float(np.percentile(ks, 97.5))) if len(ks) > 100 else (np.nan, np.nan)


def read_second(path):
    if path.lower().endswith(".csv"):
        return list(csv.DictReader(open(path))), ""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["coding"]
    coder = ws["B1"].value or ""
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and "id" in r and "c1_mt" in r)
    head = [str(h) if h is not None else "" for h in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        if r and r[head.index("id")]:
            out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(head, r)})
    return out, coder


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    first = {r["id"]: r for r in csv.DictReader(open(os.path.join(HERE, "coding.csv")))}
    second, coder = read_second(sys.argv[1])
    done = [r for r in second if r["id"] in first and all(r.get(k) for k in ITEMS)]
    partial = [r["id"] for r in second if r["id"] in first and any(r.get(k) for k in ITEMS) and r not in done]
    if not done:
        sys.exit("no completed rows in the second coder's file")
    ids = [r["id"] for r in done]
    S = {r["id"]: r for r in done}
    energy = [i for i in ids if c4_class(first[i]["c4_energy"]) != "none"]
    res = []
    for label, item, f, subset in BINARY:
        use = energy if subset == "energy" else ids
        a = [bool(f(first[i][item])) for i in use]
        b = [bool(f(S[i][item])) for i in use]
        k, po = kappa(a, b)
        lo, hi = boot_kappa(a, b)
        res.append(dict(item=label, variable=item, n=len(use), agreement_pct=round(100 * po, 1),
                        kappa=None if not np.isfinite(k) else round(k, 3), kappa_lo=None if not np.isfinite(lo) else round(lo, 3),
                        kappa_hi=None if not np.isfinite(hi) else round(hi, 3), pabak=round(2 * po - 1, 3),
                        count_coder1=int(sum(a)), count_coder2=int(sum(b))))
    fullk = {}
    for item in ITEMS:
        a = [full(item, first[i][item]) for i in ids]
        b = [full(item, S[i][item]) for i in ids]
        k, po = kappa(a, b)
        fullk[item] = dict(kappa=None if not np.isfinite(k) else round(k, 3), agreement_pct=round(100 * po, 1))
    # pooled kappa over all binary decisions (one number for the abstract of the check)
    A, B = [], []
    for label, item, f, subset in BINARY:
        use = energy if subset == "energy" else ids
        A += [bool(f(first[i][item])) for i in use]
        B += [bool(f(S[i][item])) for i in use]
    kp, pp = kappa(A, B)
    with open(os.path.join(HERE, "kappa_summary.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(res[0]))
        w.writeheader()
        w.writerows(res)
    dis = []
    for i in ids:
        for item in ITEMS:
            if full(item, first[i][item]) != full(item, S[i][item]):
                dis.append(dict(id=i, paper=first[i]["short"], item=item, coder1=first[i][item], coder2=S[i][item],
                                coder1_quotes=first[i]["evidence"], coder2_quote=S[i].get(item + "_quote", ""),
                                resolution="", final_code=""))
    with open(os.path.join(HERE, "disagreements.csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(dis[0]) if dis else ["id"])
        w.writeheader()
        w.writerows(dis)
    summ = dict(second_coder=coder, papers_compared=len(ids), papers_incomplete=partial, energy_papers=len(energy),
                pooled_binary_kappa=round(kp, 3), pooled_binary_agreement_pct=round(100 * pp, 1),
                items=res, full_code_kappa=fullk, disagreements=len(dis))
    json.dump(summ, open(os.path.join(HERE, "kappa_summary.json"), "w"), indent=1)
    print(f"second coder: {coder or '(name not filled in)'}  papers compared: {len(ids)}  (energy papers: {len(energy)})")
    print(f"{'item':44s} {'n':>3s} {'agree%':>7s} {'kappa':>6s} {'95% CI':>15s} {'PABAK':>6s} {'c1':>3s} {'c2':>3s}")
    for r in res:
        ci = f"[{r['kappa_lo']}, {r['kappa_hi']}]" if r["kappa_lo"] is not None else "-"
        print(f"{r['item']:44s} {r['n']:3d} {r['agreement_pct']:7.1f} {str(r['kappa']):>6s} {ci:>15s} "
              f"{r['pabak']:6.2f} {r['count_coder1']:3d} {r['count_coder2']:3d}")
    print(f"pooled binary kappa {kp:.3f} (agreement {100 * pp:.1f}%); {len(dis)} coded disagreements -> disagreements.csv")


if __name__ == "__main__":
    main()
