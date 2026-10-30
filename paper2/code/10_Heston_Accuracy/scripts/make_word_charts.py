#!/usr/bin/env python3
"""make_word_charts.py - the manuscript's figures as native, editable Excel charts for Word.

For every chart this writes, into OUTDIR:
  <key>.xml   a DrawingML chart part with cached values (Word draws it without opening Excel)
  <key>.xlsx  the chart's data workbook, embedded in the .docx ("Edit Data" in Word opens it)
and charts.json with each chart's size in cm. embed_word_charts.py puts them into the .docx.
Usage: python3 scripts/make_word_charts.py OUTDIR
"""
import csv, io, json, math, os, re, sys, zipfile
from copy import deepcopy
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import LineChart, ScatterChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_figures as MF

RES = MF.RES
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RES, "word_charts")
os.makedirs(OUT, exist_ok=True)

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
C, A = "{%s}" % NS["c"], "{%s}" % NS["a"]
FONTNAME = "Times New Roman"
INK, INK2, GRID = "2B2B29", "6B6A64", "D9D9D9"
COL = {"card": "2A78D6", "cpu_matched": "EB6834", "cpu_native": "1BAF7A", "gpu": "8A5CC2"}
CASE_COL = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "8A5CC2"]
PLAT_LAB = {"card": "Card (U50)", "cpu_matched": "CPU, card-matched code", "cpu_native": "CPU, native code"}
TARGETS = MF.TARGETS
LOGFMT = '[<0.01]0E+0;[>=10000]0E+0;General'
SIZES = {}
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", fgColor="E8E6E0")


# ------------------------------------------------------------------ helpers (openpyxl side)
def write_rows(ws, r0, rows, header=True):
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r0 + i, column=1 + j, value=v)
            c.font = BOLD if (header and i == 0) else Font(name="Arial", size=10)
            if header and i == 0:
                c.fill = HEAD
                c.alignment = Alignment(wrap_text=True, vertical="top")
            elif isinstance(v, float):
                c.number_format = "0.000E+00" if (abs(v) < 0.01 or abs(v) >= 1e4) and v != 0 else "0.0000"
    return r0 + len(rows)


def ln(color, w_pt=1.25, dash=None):
    lp = LineProperties(solidFill=color, w=int(w_pt * 12700))
    if dash: lp.prstDash = dash
    return lp


def series_style(s, color, marker="circle", size=5, filled=True, line_pt=None, dash=None):
    s.smooth = False
    s.marker = Marker(symbol=marker, size=size)
    s.marker.graphicalProperties = GraphicalProperties(solidFill=color if filled else "FFFFFF", ln=ln(color, 0.75))
    if line_pt:
        s.graphicalProperties = GraphicalProperties(ln=ln(color, line_pt, dash))
    else:
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def ref_line(s, color=INK2, dash="dash", w=0.75):
    s.smooth = False
    s.marker = Marker(symbol="none")
    s.graphicalProperties = GraphicalProperties(ln=ln(color, w, dash))


def style_axes(ch, xt, yt, xlog=False, ylog=False, xr=None, yr=None, xunit=None, yunit=None, legend="b"):
    for ax, title, lg, rng, unit in ((ch.x_axis, xt, xlog, xr, xunit), (ch.y_axis, yt, ylog, yr, yunit)):
        ax.title = title
        ax.delete = False
        ax.crosses = "min"
        ax.tickLblPos = "low"
        ax.majorTickMark = "out"
        ax.minorTickMark = "none"
        ax.spPr = GraphicalProperties(ln=ln("808080", 0.75))
        if lg:
            ax.scaling.logBase = 10
            ax.number_format = LOGFMT
        if rng:
            ax.scaling.min, ax.scaling.max = rng
        if unit:
            ax.majorUnit = unit
    ch.x_axis.axPos = "b"
    ch.y_axis.axPos = "l"
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=ln(GRID, 0.5)))
    ch.x_axis.majorGridlines = None
    if legend:
        ch.legend.position = legend
        ch.legend.overlay = False
    else:
        ch.legend = None


# ------------------------------------------------------------------ helpers (chart XML side)
def parse_ref(f):
    m = re.match(r"^'?(.+?)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", f)
    sh, c1, r1, c2, r2 = m.groups()
    c2, r2 = c2 or c1, r2 or r1
    return sh, column_index_from_string(c1), int(r1), column_index_from_string(c2), int(r2)


def ref_values(wb, f):
    sh, c1, r1, c2, r2 = parse_ref(f)
    ws = wb[sh]
    return [ws.cell(row=r, column=c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def el(tag, **attrs):
    e = etree.Element(tag)
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def sub(parent, tag, **attrs):
    e = etree.SubElement(parent, tag)
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def add_cache(ref_el, vals):
    """turn c:numRef/c:strRef into the right kind and append its cache"""
    is_str = any(isinstance(v, str) for v in vals)
    for old in ref_el.findall(C + "numCache") + ref_el.findall(C + "strCache"):
        ref_el.remove(old)
    ref_el.tag = C + ("strRef" if is_str else "numRef")
    cache = sub(ref_el, C + ("strCache" if is_str else "numCache"))
    if not is_str:
        sub(cache, C + "formatCode").text = "General"
    sub(cache, C + "ptCount", val=len(vals))
    for i, v in enumerate(vals):
        if v is None:
            continue
        pt = sub(cache, C + "pt", idx=i)
        sub(pt, C + "v").text = str(v) if is_str else repr(float(v))


def txpr(sz, bold=False, color=INK, rot=None):
    t = el(C + "txPr")
    bp = sub(t, A + "bodyPr")
    if rot is not None:
        bp.set("rot", str(rot)); bp.set("vert", "horz")
    sub(t, A + "lstStyle")
    p = sub(t, A + "p")
    ppr = sub(p, A + "pPr")
    d = sub(ppr, A + "defRPr", sz=sz, b=int(bold))
    sub(sub(d, A + "solidFill"), A + "srgbClr", val=color)
    sub(d, A + "latin", typeface=FONTNAME)
    sub(d, A + "cs", typeface=FONTNAME)
    sub(p, A + "endParaRPr", lang="en-US")
    return t


def rich_title(text, sz):
    t = el(C + "title")
    tx = sub(t, C + "tx")
    rich = sub(tx, C + "rich")
    sub(rich, A + "bodyPr")
    sub(rich, A + "lstStyle")
    lines = text.split("\n")
    for line in lines:
        p = sub(rich, A + "p")
        ppr = sub(p, A + "pPr")
        d = sub(ppr, A + "defRPr", sz=sz, b=0)
        sub(d, A + "latin", typeface=FONTNAME)
        r = sub(p, A + "r")
        rp = sub(r, A + "rPr", lang="en-US", sz=sz, b=0)
        sub(sub(rp, A + "solidFill"), A + "srgbClr", val=INK)
        sub(rp, A + "latin", typeface=FONTNAME)
        sub(r, A + "t").text = line
    sub(t, C + "overlay", val=0)
    return t


def insert_ordered(parent, new, after_tags):
    """insert new after the last existing child whose tag is in after_tags (schema order), else first"""
    idx = 0
    for i, ch in enumerate(parent):
        if ch.tag in after_tags:
            idx = i + 1
    parent.insert(idx, new)


def patch(xml, wb, title, catrot=None):
    root = etree.fromstring(xml)
    # caches for every reference
    for ref in root.iter(C + "numRef", C + "strRef"):
        add_cache(ref, ref_values(wb, ref.find(C + "f").text))
    chart = root.find(C + "chart")
    # chart title
    for t in chart.findall(C + "title"):
        chart.remove(t)
    for t in chart.findall(C + "autoTitleDeleted"):
        chart.remove(t)
    if title:
        chart.insert(0, rich_title(title, 900))
        chart.insert(1, el(C + "autoTitleDeleted", val=0))
    else:
        chart.insert(0, el(C + "autoTitleDeleted", val=1))
    # axis titles: plain 8 pt
    for ax in chart.iter(C + "valAx", C + "catAx"):
        t = ax.find(C + "title")
        if t is not None:
            text = "\n".join("".join(x.text or "" for x in p.iter(A + "t")) for p in t.iter(A + "p"))
            new = rich_title(text, 800)
            ax.replace(t, new)
            if ax.tag == C + "valAx" and ax.find(C + "axPos").get("val") in ("l", "r"):
                body = new.find(".//" + A + "bodyPr")
                body.set("rot", "-5400000")
                body.set("vert", "horz")
        # tick labels 8 pt
        for old in ax.findall(C + "txPr"):
            ax.remove(old)
        rot = catrot if (catrot is not None and ax.tag == C + "catAx") else None
        insert_ordered(ax, txpr(800, color=INK2, rot=rot), {C + x for x in ("axId", "scaling", "delete", "axPos", "majorGridlines",
                       "minorGridlines", "title", "numFmt", "majorTickMark", "minorTickMark", "tickLblPos", "spPr")})
    leg = chart.find(C + "legend")
    if leg is not None:
        for old in leg.findall(C + "txPr"):
            leg.remove(old)
        insert_ordered(leg, txpr(750), {C + x for x in ("legendPos", "legendEntry", "layout", "overlay", "spPr")})
    if chart.find(C + "plotVisOnly") is None:
        insert_ordered(chart, el(C + "plotVisOnly", val=1), {C + x for x in ("title", "autoTitleDeleted", "pivotFmts", "view3D",
                       "floor", "sideWall", "backWall", "plotArea", "legend")})
    if chart.find(C + "dispBlanksAs") is None:
        insert_ordered(chart, el(C + "dispBlanksAs", val="gap"), {C + x for x in ("title", "autoTitleDeleted", "pivotFmts", "view3D",
                       "floor", "sideWall", "backWall", "plotArea", "legend", "plotVisOnly")})
    # chart space: square corners, no border, fonts, link to the embedded workbook
    for tag in ("roundedCorners", "spPr", "txPr", "externalData"):
        for old in root.findall(C + tag):
            root.remove(old)
    insert_ordered(root, el(C + "roundedCorners", val=0), {C + "date1904", C + "lang"})
    sp = el(C + "spPr")
    sub(sp, A + "noFill")
    sub(sub(sp, A + "ln"), A + "noFill")
    insert_ordered(root, sp, {C + x for x in ("date1904", "lang", "roundedCorners", "AlternateContent", "style", "clrMapOvr",
                   "pivotSource", "protection", "chart")})
    insert_ordered(root, txpr(800), {C + x for x in ("date1904", "lang", "roundedCorners", "style", "clrMapOvr", "pivotSource",
                   "protection", "chart", "spPr")})
    ext = el(C + "externalData")
    ext.set("{%s}id" % NS["r"], "rId1")
    sub(ext, C + "autoUpdate", val=0)
    insert_ordered(root, ext, {C + x for x in ("date1904", "lang", "roundedCorners", "style", "clrMapOvr", "pivotSource",
                   "protection", "chart", "spPr", "txPr")})
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


MASTER = []          # (key, data rows, chart, notes) for the stand-alone figure workbook


def emit(key, wb, ws, chart, title, w_cm, h_cm, notes, catrot=None):
    MASTER.append((key, [[c.value for c in row] for row in ws.iter_rows()],
                   {k: v.width for k, v in ws.column_dimensions.items()}, deepcopy(chart), title, w_cm, h_cm, list(notes)))
    src = wb.create_sheet("Source")
    for i, n in enumerate(notes, 1):
        src.cell(row=i, column=1, value=n).font = Font(name="Arial", size=10, italic=(i > 1), bold=(i == 1))
    src.column_dimensions["A"].width = 140
    chart.width, chart.height = w_cm, h_cm
    ws.add_chart(chart, "R2")
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        name = [n for n in z.namelist() if n.startswith("xl/charts/chart")][0]
        xml = z.read(name)
    open(os.path.join(OUT, key + ".xml"), "wb").write(patch(xml, wb, title, catrot))
    ws._charts = []                       # the embedded workbook holds the data only; the chart lives in Word
    ws.sheet_view.zoomScale = 100
    wb.save(os.path.join(OUT, key + ".xlsx"))
    SIZES[key] = [w_cm, h_cm]


# ------------------------------------------------------------------ Part A data
rows = MF.load()
G = {(x["case"], x["target"], x["boundary"], x["combo"]): x for x in rows}
SRC_A = "Source: results/analysis/measured_vs_predicted.csv (main cases) and results/analysis_s/measured_vs_predicted_s.csv (secondary cases, variant 'same'); medians of two sessions x 5 repeats."

# Figs 1-6: one line chart per panel. Series order sets the drawing order: card last (on top) with the smallest
# marker and line, CPU-native behind it with the largest, so coinciding values stay visible.
PANEL_W, PANEL_H = 8.2, 8.0
SHORT = {"card": "Card", "cpu_matched": "CPU matched", "cpu_native": "CPU native"}
MSHORT = {"ft": "FT", "qe": "QE", "plain": "plain", "cv": "CV"}
for n, (case, name) in enumerate(MF.CASES, 1):
    a, b = MF.METHODS[case]
    order = [("cpu_matched", a), ("cpu_matched", b), ("cpu_native", a), ("cpu_native", b), ("card", a), ("card", b)]
    look = {"cpu_matched": (5, 1.25), "cpu_native": (7, 2.25), "card": (4, 1.0)}
    for k, (field, title) in enumerate((("E_meas_J", "(a) One pricing run"), ("Ess_meas_J", "(b) Steady state"))):
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        r = 1
        blocks = {}
        for fld, lab in (("E_meas_J", "(a) One pricing run: energy to target (J)"), ("Ess_meas_J", "(b) Steady state: energy to target (J)")):
            ws.cell(row=r, column=1, value=lab).font = BOLD
            hdr = ["RMSE target"] + [f"{SHORT[p]}, {MSHORT[m]}" for p, m in order]
            body = [[t] + [G[(case, t, "socket", f"{p}_{m}")][fld] for p, m in order] for t in TARGETS]
            blocks[fld] = r + 1
            r = write_rows(ws, r + 1, [hdr] + body) + 1
        for j in range(1, 8):
            ws.column_dimensions[get_column_letter(j)].width = 18 if j > 1 else 12
        h = blocks[field]
        ch = LineChart()
        for j, (p, m) in enumerate(order):
            s = Series(Reference(ws, min_col=2 + j, min_row=h, max_row=h + 3), title_from_data=True)
            size, lw = look[p]
            series_style(s, COL[p], marker="circle" if m == a else "square", size=size, filled=(m == a), line_pt=lw,
                         dash=None if m == a else "dash")
            ch.series.append(s)
        ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=h + 3))
        vals = [G[(case, t, "socket", f"{p}_{m}")][field] for t in TARGETS for p, m in order]
        style_axes(ch, "RMSE target", "Energy to target (J)", ylog=True,
                   yr=(10 ** math.floor(math.log10(min(vals))), 10 ** math.ceil(math.log10(max(vals)))), legend="b")
        emit(f"case{case}_{'ab'[k]}", wb, ws, ch, title, PANEL_W, PANEL_H,
             [f"Fig. {n}{'ab'[k]}: energy to target, {name}, socket boundary", SRC_A,
              "Edit the values on the Data sheet; the chart in the manuscript follows them."])

# Fig 7: method gain against platform gain; marker shape by target
pts = MF.fig2(rows)
wb = Workbook(); ws = wb.active; ws.title = "Data"
hdr = ["Case", "Target", "Gain from the better platform", "Gain from the better method"]
body = [[c, t, x, y] for (c, t, x, y) in pts]
end = write_rows(ws, 1, [hdr] + body)
write_rows(ws, end + 1, [["Equal-gain line", "", "x", "y"], ["", "", 1.0, 1.0], ["", "", 2.6, 2.6]])
for j, w in enumerate((10, 8, 22, 22), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ch = ScatterChart()
ch.scatterStyle = "lineMarker"
shape = {"1%": "circle", "0.3%": "square", "0.1%": "triangle"}
for i, ((case, _), c) in enumerate(zip(MF.CASES, CASE_COL)):
    r0 = 2 + 3 * i
    s = Series(Reference(ws, min_col=4, min_row=r0, max_row=r0 + 2), Reference(ws, min_col=3, min_row=r0, max_row=r0 + 2),
               title=f"Case {case}")
    series_style(s, c, marker="circle", size=6, line_pt=0.75)
    s.dPt = []
    for k, t in enumerate(TARGETS):
        dp = DataPoint(idx=k)
        dp.marker = Marker(symbol=shape[t], size=6)
        dp.marker.graphicalProperties = GraphicalProperties(solidFill=c, ln=ln("FFFFFF", 0.5))
        s.dPt.append(dp)
    ch.series.append(s)
s = Series(Reference(ws, min_col=4, min_row=end + 2, max_row=end + 3), Reference(ws, min_col=3, min_row=end + 2, max_row=end + 3),
           title="equal gain")
ref_line(s, dash="sysDot")
ch.series.append(s)
style_axes(ch, "Gain from the better platform (steady state, socket)", "Gain from the better method", xlog=False, ylog=True,
           xr=(1, 2.5), yr=(1, 1000), xunit=0.5, legend="r")
emit("gains", wb, ws, ch, None, 10.5, 8.5, ["Fig. 7: method gain against platform gain (steady state, socket boundary)",
     "Derived from the steady-state energies of Figs 1-6 (make_figures.py fig2). Marker: circle 1%, square 0.3%, triangle 0.1%."])

# Figs 8 and 9: prediction
for key, pc, mc, title, n in (("pred_one", "E_pred_J", "E_meas_J", "One pricing run", 8), ("pred_ss", "Ess_pred_J", "Ess_meas_J", "Steady state", 9)):
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    hdr = ["Platform", "Case", "Target", "Boundary", "Combination", "Predicted energy (J)", "Measured energy (J)"]
    body, spans, r = [], {}, 2
    for p in ("cpu_native", "cpu_matched", "card"):
        sel = [x for x in rows if x["combo"].startswith(p)]
        spans[p] = (r, r + len(sel) - 1)
        r += len(sel)
        body += [[PLAT_LAB[p], x["case"], x["target"], x["boundary"], x["combo"], x[pc], x[mc]] for x in sel]
    end = write_rows(ws, 1, [hdr] + body)
    write_rows(ws, end + 1, [["Reference lines", "", "", "", "", "x (J)", "y (J)"]])
    lines = {"measured = predicted": 1.0, "+20%": 1.2, "−20%": 1 / 1.2}
    rr = end + 2
    line_rows = {}
    for nm, f in lines.items():
        write_rows(ws, rr, [[nm, "", "", "", "", 1e-4, 1e-4 * f], ["", "", "", "", "", 1e5, 1e5 * f]], header=False)
        line_rows[nm] = rr
        rr += 2
    for j, w in enumerate((22, 6, 7, 9, 14, 18, 18), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ch = ScatterChart()
    ch.scatterStyle = "lineMarker"
    for p, mk in (("cpu_native", "triangle"), ("cpu_matched", "square"), ("card", "circle")):
        a0, a1 = spans[p]
        s = Series(Reference(ws, min_col=7, min_row=a0, max_row=a1), Reference(ws, min_col=6, min_row=a0, max_row=a1), title=PLAT_LAB[p])
        series_style(s, COL[p], marker=mk, size=4)
        ch.series.append(s)
    for nm, r0 in line_rows.items():
        s = Series(Reference(ws, min_col=7, min_row=r0, max_row=r0 + 1), Reference(ws, min_col=6, min_row=r0, max_row=r0 + 1), title=nm)
        ref_line(s, color=INK2, dash=None if nm.startswith("measured") else "dash", w=0.75)
        ch.series.append(s)
    style_axes(ch, "Predicted energy (J)", "Measured energy (J)", xlog=True, ylog=True, xr=(1e-4, 1e5), yr=(1e-4, 1e5), legend="b")
    emit(key, wb, ws, ch, title, 9.5, 10.0, [f"Fig. {n}: measured against predicted energy ({title.lower()}), every combination, socket and server boundaries",
         SRC_A + " Predictions: registered files (secondary cases, variant 'same')."])

# ------------------------------------------------------------------ Part B
d = json.load(open(os.path.join(RES, "analysis_gpu_cpu48", "analysis_gpu.json")))
wb = Workbook(); ws = wb.active; ws.title = "Data"
hdr = ["Case and target", "U50 card", "CPU, native code (48 threads)", "V100 GPU", "Card, first method", "Card, second method",
       "CPU-native, first", "CPU-native, second", "V100, first", "V100, second"]
hdr = hdr[:4] + ["RTX 4070 Ti SUPER"] + hdr[4:] + ["RTX, first", "RTX, second"]
RTXA = json.load(open(os.path.join(RES, "analysis_rtx_v09", "rtx_v09.json")))
RTXE = {(x["case"], x["target"], x["scheme"]): x["Ess_meas_J"] for x in RTXA["per_setting"]}
body = []
for w in d["winners"]:
    if w["boundary"] != "socket":
        continue
    m1, m2 = MF.METHODS[w["case"]]
    e = [w["Ess"][k] for k in (f"card_{m1}", f"card_{m2}", f"cpu_native_{m1}", f"cpu_native_{m2}", f"gpu_{m1}", f"gpu_{m2}")]
    r1, r2 = RTXE[(w["case"], w["target"], m1)], RTXE[(w["case"], w["target"], m2)]
    body.append([f"{w['case']} {w['target']}", min(e[0:2]), min(e[2:4]), min(e[4:6]), min(r1, r2)] + e + [r1, r2])
end = write_rows(ws, 1, [hdr] + body)
for j in range(1, 13):
    ws.column_dimensions[get_column_letter(j)].width = 16
ch = LineChart()
COL["rtx"] = "EDA100"
for col, p, mk, sz in ((3, "cpu_native", "square", 7), (2, "card", "circle", 5), (4, "gpu", "diamond", 6), (5, "rtx", "triangle", 6)):
    s = Series(Reference(ws, min_col=col, min_row=1, max_row=end - 1), title_from_data=True)
    series_style(s, COL[p], marker=mk, size=sz)
    ch.series.append(s)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=end - 1))
style_axes(ch, "Case and target", "Energy per option, steady state (J)", ylog=True, yr=(1e-5, 1000), legend="t")
emit("three", wb, ws, ch, None, 16.5, 8.0, ["Steady-state energy per option, lower-energy method on each platform, primary boundaries",
     "Source: results/analysis_gpu_cpu48/analysis_gpu.json (winners, Ess, boundary 'socket') and results/analysis_rtx_v09/rtx_v09.json "
     "(RTX: GPU + CPU package, medians of two sessions x 5 repeats). Columns B-E are the lower of the two methods; "
     "they are plain values here so that the chart can be edited directly."], catrot=-2700000)

# Fig 11: co-execution
def coexec(fn):
    R = {}
    for x in csv.DictReader(open(os.path.join(RES, fn))):
        if x["workload"] == "large":
            R[(x["case"], x["scheme"], x["policy"])] = x
    out = {}
    for (c, sch, pol) in list(R):
        if pol == "acc":
            t = R.get((c, sch, "static_T")) or R.get((c, sch, "static_E"))
            a = R[(c, sch, "acc")]
            out[(c, sch)] = (float(t["makespan_median_s"]) / float(a["makespan_median_s"]), float(t["E_primary_J"]) / float(a["E_primary_J"]))
    return out


F_, S_ = coexec("coexec_frodo_summary.csv"), coexec("coexec_smc_summary.csv")
wb = Workbook(); ws = wb.active; ws.title = "Data"
hdr = ["Kernel", "frodo: makespan ratio", "frodo: energy ratio", "GPU server: makespan ratio", "GPU server: energy ratio"]
body = [[f"{c} {m.upper() if m in ('ft', 'qe') else m}", F_[(c, m)][0], F_[(c, m)][1], S_[(c, m)][0], S_[(c, m)][1]]
        for c, _ in MF.CASES for m in MF.METHODS[c]]
end = write_rows(ws, 1, [hdr] + body)
write_rows(ws, end + 1, [["Reference lines", "x", "y", "x", "y"], ["equal energy", 0.2, 1.0, 1.0, 0.6], ["", 1.1, 1.0, 1.0, 1.3]])
for j, w in enumerate((14, 18, 18, 22, 22), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ch = ScatterChart()
ch.scatterStyle = "lineMarker"
for xc, yc, t, p, mk in ((2, 3, "U50 + 46 CPU threads (frodo)", "card", "circle"), (4, 5, "V100 + 18 CPU threads (GPU server)", "gpu", "diamond")):
    s = Series(Reference(ws, min_col=yc, min_row=2, max_row=end - 1), Reference(ws, min_col=xc, min_row=2, max_row=end - 1), title=t)
    series_style(s, COL[p], marker=mk, size=6)
    ch.series.append(s)
for xc, yc, t in ((2, 3, "equal energy"), (4, 5, "equal makespan")):
    s = Series(Reference(ws, min_col=yc, min_row=end + 2, max_row=end + 3), Reference(ws, min_col=xc, min_row=end + 2, max_row=end + 3), title=t)
    ref_line(s, dash="sysDot")
    ch.series.append(s)
style_axes(ch, "Makespan, time-optimal split / accelerator alone", "Energy, time-optimal split / accelerator alone",
           xr=(0.2, 1.1), yr=(0.6, 1.3), xunit=0.1, yunit=0.1, legend="b")
ch.x_axis.number_format = "0.0"; ch.y_axis.number_format = "0.0"
emit("coexec", wb, ws, ch, None, 10.0, 9.0, ["Fig. 11: co-execution, static_T against the accelerator alone, large workloads, primary boundary",
     "Source: results/coexec_frodo_summary.csv and results/coexec_smc_summary.csv (median makespan and energy per run, 3 repeats). "
     "Where static_T had the same CPU fraction as static_E it was run once, under the name static_E."])

# Fig 12: scheduling
POL = ["alone_card", "alone_fcpu", "alone_scpu", "round_robin", "alone_gpu", "model_dev", "model_E", "model_M", "model_T", "oracle"]
NAME = {"alone_card": "card alone", "alone_fcpu": "frodo CPU alone", "alone_scpu": "GPU-server CPU alone", "round_robin": "round robin",
        "alone_gpu": "V100 alone", "model_dev": "model_dev", "model_E": "model_E", "model_M": "model_M", "model_T": "model_T", "oracle": "oracle"}
PC = {"alone_card": COL["card"], "alone_fcpu": COL["cpu_native"], "alone_scpu": COL["cpu_matched"], "round_robin": INK2,
      "alone_gpu": COL["gpu"], "model_dev": INK2, "model_E": INK, "model_M": INK, "model_T": INK, "oracle": INK}
PM = {"alone_card": "circle", "alone_fcpu": "square", "alone_scpu": "triangle", "round_robin": "x", "alone_gpu": "diamond",
      "model_dev": "triangle", "model_E": "star", "model_M": "plus", "model_T": "dash", "oracle": "circle"}
raw = list(csv.DictReader(open(os.path.join(RES, "sched_summary.csv"))))
FAST = ("alone_gpu", "model_dev", "model_E", "model_M", "model_T", "oracle")
for key, title, pols, xlog in (("sched_a", "(a) All policies", POL, True), ("sched_b", "(b) Fast policies", [p for p in POL if p in FAST], False)):
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    body, spans, r = [], {}, 2
    for p in POL:
        sel = sorted([x for x in raw if x["policy"] == p], key=lambda x: int(x["rep"]))
        spans[p] = (r, r + len(sel) - 1)
        r += len(sel)
        body += [[NAME[p], int(x["rep"]), float(x["T_pass_s"]), float(x["E_pass_J"])] for x in sel]
    write_rows(ws, 1, [["Policy", "Repeat", "Time per pass (s)", "Energy per pass (J)"]] + body)
    for j, w in enumerate((22, 8, 18, 18), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ch = ScatterChart()
    ch.scatterStyle = "lineMarker"
    for p in pols:
        a0, a1 = spans[p]
        s = Series(Reference(ws, min_col=4, min_row=a0, max_row=a1), Reference(ws, min_col=3, min_row=a0, max_row=a1), title=NAME[p])
        series_style(s, PC[p], marker=PM[p], size=6 if PM[p] in ("x", "plus", "star", "dash") else 5)
        if PM[p] in ("x", "plus", "star", "dash"):
            s.marker.graphicalProperties = GraphicalProperties(solidFill=PC[p], ln=ln(PC[p], 1.0))
        ch.series.append(s)
    if xlog:
        style_axes(ch, "Time per pass of the 120-option book (s)", "Energy per pass, two servers (J)", xlog=True, ylog=True,
                   xr=(0.1, 100), yr=(10, 10000), legend="b")
    else:
        style_axes(ch, "Time per pass (s)", "Energy per pass (J)", xr=(0.165, 0.21), yr=(48, 66), xunit=0.01, yunit=2, legend="b")
        ch.x_axis.number_format = "0.00"; ch.y_axis.number_format = "0"
    emit(key, wb, ws, ch, title, 8.2, 9.5, [f"Fig. 12{key[-1]}: scheduling a 120-option book on two servers, every repeat",
         "Source: results/sched_summary.csv (5 repeats x 10 policies, two-server boundary)."])

# ------------------------------------------------------------------ addendum v0.9: the RTX 4070 Ti SUPER
per = RTXA["per_setting"]
wb = Workbook(); ws = wb.active; ws.title = "Data"
hdr = ["Case", "Target", "Method", "Binding pipe (predicted)", "Predicted energy (J)", "Measured energy (J)",
       "Time per path-step, measured / predicted", "GPU power, measured / predicted"]
fp = [x for x in per if x["bind_ss"] == "fp64"]
ot = [x for x in per if x["bind_ss"] != "fp64"]
body = [[x["case"], x["target"], x["scheme"].upper() if x["scheme"] in ("ft", "qe") else x["scheme"], x["bind_ss"],
         x["Ess_pred_J"], x["Ess_meas_J"], x["t_ratio"], x["P_ratio"]] for x in ot + fp]
end = write_rows(ws, 1, [hdr] + body)
lines = {"measured = predicted": 1.0, "+30%": 1.3, "−30%": 1 / 1.3}
rr, line_rows = end + 1, {}
write_rows(ws, rr, [["Reference lines", "", "", "", "x (J)", "y (J)"]]); rr += 1
for nm, f in lines.items():
    write_rows(ws, rr, [[nm, "", "", "", 1e-5, 1e-5 * f], ["", "", "", "", 1e3, 1e3 * f]], header=False)
    line_rows[nm] = rr; rr += 2
for j, w in enumerate((7, 7, 8, 12, 16, 16, 18, 18), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
n_o, n_f = len(ot), len(fp)
ch = ScatterChart(); ch.scatterStyle = "lineMarker"
for (a0, a1, t, c, mk) in ((2, 1 + n_o, "issue-, latency- or xu-bound", COL["gpu"], "circle"),
                           (2 + n_o, 1 + n_o + n_f, "fp64-bound", COL["cpu_matched"], "triangle")):
    s = Series(Reference(ws, min_col=6, min_row=a0, max_row=a1), Reference(ws, min_col=5, min_row=a0, max_row=a1), title=t)
    series_style(s, c, marker=mk, size=5); ch.series.append(s)
for nm, r0 in line_rows.items():
    s = Series(Reference(ws, min_col=6, min_row=r0, max_row=r0 + 1), Reference(ws, min_col=5, min_row=r0, max_row=r0 + 1), title=nm)
    ref_line(s, color=INK2, dash=None if nm.startswith("measured") else "dash", w=0.75); ch.series.append(s)
style_axes(ch, "Predicted energy per option (J)", "Measured energy per option (J)", xlog=True, ylog=True, xr=(1e-5, 1e3), yr=(1e-5, 1e3), legend="b")
emit("rtx_e", wb, ws, ch, "(a) Energy per option", 8.2, 9.5,
     ["RTX 4070 Ti SUPER: measured against registered predicted steady-state energy per option (GPU + CPU package), 36 settings",
      "Source: results/analysis_rtx_v09/rtx_v09.json (per_setting); prediction SHA-256 4bee4fe5...; medians of two sessions x 5 repeats."])

wb = Workbook(); ws = wb.active; ws.title = "Data"
end = write_rows(ws, 1, [hdr] + body)
curve = [[x, 1 / x] for x in (0.9, 1.0, 1.2, 1.4, 1.6, 1.8, 2.0, 2.2)]
write_rows(ws, end + 1, [["Equal energy (GPU only): power ratio = 1 / time ratio", "", "", "", "", "", "x", "y"]] +
           [["", "", "", "", "", "", a, b] for a, b in curve], header=False)
for j, w in enumerate((7, 7, 8, 12, 16, 16, 18, 18), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ch = ScatterChart(); ch.scatterStyle = "lineMarker"
for (a0, a1, t, c, mk) in ((2, 1 + n_o, "issue-, latency- or xu-bound", COL["gpu"], "circle"),
                           (2 + n_o, 1 + n_o + n_f, "fp64-bound", COL["cpu_matched"], "triangle")):
    s = Series(Reference(ws, min_col=8, min_row=a0, max_row=a1), Reference(ws, min_col=7, min_row=a0, max_row=a1), title=t)
    series_style(s, c, marker=mk, size=5); ch.series.append(s)
s = Series(Reference(ws, min_col=8, min_row=end + 2, max_row=end + 1 + len(curve)), Reference(ws, min_col=7, min_row=end + 2, max_row=end + 1 + len(curve)),
           title="equal GPU energy")
ref_line(s, dash="sysDot"); ch.series.append(s)
style_axes(ch, "Time per path-step, measured / predicted", "GPU power, measured / predicted", xr=(0.9, 2.2), yr=(0.4, 1.4), xunit=0.2, yunit=0.2, legend="b")
ch.x_axis.number_format = "0.0"; ch.y_axis.number_format = "0.0"
emit("rtx_tp", wb, ws, ch, "(b) Rate and power", 8.2, 9.5,
     ["RTX 4070 Ti SUPER: measured / predicted time per path-step and GPU power, 36 settings",
      "Source: results/analysis_rtx_v09/rtx_v09.json (per_setting: t_ratio, P_ratio)."])

# ------------------------------------------------------------------ addendum v0.9: the corrected scheduler
POL9 = ["alone_card", "alone_fcpu", "alone_scpu", "alone_gpu", "model_dev", "model_E", "model_M", "model_T", "oracle", "oracle_M"]
NAME.update({"oracle_M": "oracle_M"}); PC.update({"oracle_M": INK2}); PM.update({"oracle_M": "diamond"})
raw9 = list(csv.DictReader(open(os.path.join(RES, "sched_v09_summary.csv"))))
FAST9 = ("alone_gpu", "model_dev", "model_E", "model_M", "model_T", "oracle", "oracle_M")
for key, title, pols, xlog in (("sched9_a", "(a) All policies", POL9, True), ("sched9_b", "(b) Fast policies", [p for p in POL9 if p in FAST9], False)):
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    body, spans, r = [], {}, 2
    for p in POL9:
        sel = sorted([x for x in raw9 if x["policy"] == p and x["regime"] in ("R1", "R2")], key=lambda x: int(x["rep"]))
        spans[p] = (r, r + len(sel) - 1)
        r += len(sel)
        body += [[NAME[p], int(x["rep"]), float(x["T_pass_s"]), float(x["E_pass_J"]), float(x["T_pred_s"]) if x["T_pred_s"] else None,
                  float(x["E_pred_J"]) if x["E_pred_J"] else None] for x in sel]
    write_rows(ws, 1, [["Policy", "Repeat", "Time per pass (s)", "Energy per pass (J)", "Predicted time (s)", "Predicted energy (J)"]] + body)
    for j, w in enumerate((22, 8, 18, 18, 18, 18), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ch = ScatterChart(); ch.scatterStyle = "lineMarker"
    for p in pols:
        a0, a1 = spans[p]
        s = Series(Reference(ws, min_col=4, min_row=a0, max_row=a1), Reference(ws, min_col=3, min_row=a0, max_row=a1), title=NAME[p])
        series_style(s, PC[p], marker=PM[p], size=6 if PM[p] in ("x", "plus", "star", "dash") else 5)
        if PM[p] in ("x", "plus", "star", "dash"):
            s.marker.graphicalProperties = GraphicalProperties(solidFill=PC[p], ln=ln(PC[p], 1.0))
        ch.series.append(s)
    if xlog:
        style_axes(ch, "Time per pass of the 120-option book (s)", "Energy per pass, two servers (J)", xlog=True, ylog=True,
                   xr=(0.1, 100), yr=(10, 10000), legend="b")
    else:
        style_axes(ch, "Time per pass (s)", "Energy per pass (J)", xr=(0.14, 0.18), yr=(44, 72), xunit=0.01, yunit=4, legend="b")
        ch.x_axis.number_format = "0.00"; ch.y_axis.number_format = "0"
    emit(key, wb, ws, ch, title, 8.2, 9.5, ["Scheduling with the corrected model (addendum v0.9), regimes R1 and R2, every repeat",
         "Source: results/sched_v09_summary.csv (5 repeats, two-server boundary; registered predictions in columns E-F)."])

# ------------------------------------------------------------------ the stand-alone figure workbook (data + native Excel charts)
if len(sys.argv) > 2:
    from openpyxl.chart.layout import Layout
    FIGNAME = json.loads(os.environ.get("FIGNAMES", "{}"))
    mw = Workbook(); idx = mw.active; idx.title = "Index"
    idx.append(["Sheet", "Figure", "What it shows", "Source"])
    for c in idx[1]:
        c.font = BOLD; c.fill = HEAD
    def _order(m):
        lab = FIGNAME.get(m[0], "Z" + m[0]).replace("Fig ", "")
        sup = lab.startswith("S")
        num = "".join(ch for ch in lab if ch.isdigit())
        return (sup, int(num) if num else 99, lab)
    for key, rows_, widths, chart, title, w_cm, h_cm, notes in sorted(MASTER, key=_order):
        name = FIGNAME.get(key, key)[:31]
        ws2 = mw.create_sheet(name.replace(".", "").replace(":", "")[:31])
        for i, row in enumerate(rows_, 1):
            for j, v in enumerate(row, 1):
                cc = ws2.cell(row=i, column=j, value=v)
                cc.font = BOLD if i == 1 else Font(name="Arial", size=10)
                if isinstance(v, float):
                    cc.number_format = "0.000E+00" if (abs(v) < 0.01 or abs(v) >= 1e4) and v != 0 else "0.0000"
        for k, wv in widths.items():
            if wv: ws2.column_dimensions[k].width = wv
        q = "'" + ws2.title + "'!"
        for srs in chart.series:
            for attr in ("val", "cat", "xVal", "yVal"):
                ref = getattr(srs, attr, None)
                if ref is None: continue
                for sub_ in ("numRef", "strRef"):
                    rr_ = getattr(ref, sub_, None)
                    if rr_ is not None and rr_.f:
                        rr_.f = rr_.f.replace("'Data'!", q).replace("Data!", q)
            if srs.tx is not None and srs.tx.strRef is not None and srs.tx.strRef.f:
                srs.tx.strRef.f = srs.tx.strRef.f.replace("'Data'!", q).replace("Data!", q)
        if title:
            chart.title = title
        chart.width, chart.height = w_cm * 1.25, h_cm * 1.25
        col = get_column_letter(max(len(r) for r in rows_) + 2)
        ws2.add_chart(chart, f"{col}2")
        idx.append([ws2.title, FIGNAME.get(key, key), notes[0], " ".join(notes[1:2])])
    for col_, w in zip("ABCD", (16, 18, 90, 120)):
        idx.column_dimensions[col_].width = w
    mw.save(sys.argv[2])
    print("figure workbook ->", sys.argv[2])

json.dump(SIZES, open(os.path.join(OUT, "charts.json"), "w"), indent=1)
print("charts:", len(SIZES), "->", OUT)
